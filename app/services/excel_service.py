"""Excel export: Pivot (as shown in the app), Data (flat) and Query (provenance) sheets."""

import io
from datetime import datetime, timezone
from typing import Any, Dict, List, Optional

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from config import settings

HEADER_FILL = PatternFill("solid", fgColor="016FB7")
HEADER_FONT = Font(bold=True, color="FFFFFF")
GROUP_FONT = Font(bold=True)

CURRENCY_HINTS = ("revpar", "adr", "revenue", "rate")
PERCENT_HINTS = ("occupancy", "occ", "_pct", "change", "share")
INDEX_HINTS = ("rgi", "mpi", "ari")


def number_format(col: str) -> Optional[str]:
    c = col.lower()
    if any(h in c for h in INDEX_HINTS):
        return "0.0"
    if any(h in c for h in CURRENCY_HINTS):
        return '"$"#,##0.00'
    if any(h in c for h in PERCENT_HINTS):
        return "0.0%"
    return None


def to_cell(v: Any) -> Any:
    """Genie returns every value as a string; store numbers as numbers so Excel can sum them."""
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return v
    s = str(v)
    try:
        return int(s) if s.lstrip("-").isdigit() else float(s)
    except ValueError:
        return s


def _style_header(ws, row: int, ncols: int):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _autofit(ws, max_width: int = 48):
    for col in ws.columns:
        width = max((len(str(c.value)) for c in col if c.value is not None), default=8)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max_width, max(10, width + 2))


def write_pivot(ws, pivot: Dict[str, Any]):
    """Write the pivot exactly as the app computed it: header row + body rows (+ optional totals)."""
    header: List[str] = pivot["header"]
    ws.append(header)
    _style_header(ws, 1, len(header))
    fmt = number_format(pivot.get("value_label", "") or ",".join(pivot.get("values", [])))
    n_row_fields = len(pivot.get("rows", [])) or 1
    for r in pivot["body"]:
        ws.append([to_cell(v) for v in r])
    if pivot.get("totals"):
        ws.append([to_cell(v) for v in pivot["totals"]])
        for c in ws[ws.max_row]:
            c.font = GROUP_FONT
    for row in ws.iter_rows(min_row=2, min_col=n_row_fields + 1):
        for c in row:
            if fmt and isinstance(c.value, (int, float)):
                c.number_format = fmt
    ws.freeze_panes = ws.cell(row=2, column=n_row_fields + 1)
    _autofit(ws)


def write_data(ws, columns: List[str], rows: List[List[Any]]):
    ws.append(columns)
    _style_header(ws, 1, len(columns))
    for r in rows:
        ws.append([to_cell(v) for v in r])
    for i, col in enumerate(columns, start=1):
        fmt = number_format(col)
        if fmt:
            for (c,) in ws.iter_rows(min_row=2, min_col=i, max_col=i):
                if isinstance(c.value, (int, float)):
                    c.number_format = fmt
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    _autofit(ws)


def write_query(ws, question: str, sql: str, agent_name: str, data_as_of: str, row_count: int, truncated: bool):
    meta = [
        ("Question", question),
        ("Genie agent", agent_name),
        ("Generated at (UTC)", datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M")),
        ("Data as of", data_as_of),
        ("Rows exported", row_count),
        ("Truncated", "Yes: Genie returned a partial result" if truncated else "No"),
        ("Source", f"Unity Catalog metric views, {settings.catalog()}.{settings.schema()} (RevSight IQ demo data)"),
        ("SQL", sql),
    ]
    for k, v in meta:
        ws.append([k, v])
        ws.cell(row=ws.max_row, column=1).font = GROUP_FONT
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 110
    for row in ws.iter_rows(min_col=2, max_col=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")


def export_to_excel(question: str, sql: str, agent_name: str, columns: List[str], rows: List[List[Any]],
                    data_as_of: str, pivot_spec: Optional[Dict[str, Any]] = None, truncated: bool = False) -> io.BytesIO:
    wb = Workbook()
    first = wb.active
    if pivot_spec and pivot_spec.get("header") and pivot_spec.get("body"):
        first.title = "Pivot"
        write_pivot(first, pivot_spec)
        data_ws = wb.create_sheet("Data")
    else:
        first.title = "Data"
        data_ws = first
    write_data(data_ws, columns, rows)
    write_query(wb.create_sheet("Query"), question, sql, agent_name, data_as_of, len(rows), truncated)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
